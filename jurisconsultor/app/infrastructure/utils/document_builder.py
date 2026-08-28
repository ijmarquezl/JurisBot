import os
import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from fpdf import FPDF

def create_docx_from_html(html_content: str, output_path: str):
    """
    Parses HTML content (generated from markdown) and creates a DOCX file.
    """
    document = Document()
    soup = BeautifulSoup(html_content, 'html.parser')

    for element in soup.body.children if soup.body else soup.children:
        if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            level = int(element.name[1])
            heading = document.add_heading(element.get_text(), level=level)
        elif element.name == 'p':
            p = document.add_paragraph()
            for child in element.children:
                if child.name == 'strong' or child.name == 'b':
                    p.add_run(child.get_text()).bold = True
                elif child.name == 'em' or child.name == 'i':
                    p.add_run(child.get_text()).italic = True
                elif child.name is None:
                    p.add_run(str(child))
                else:
                    p.add_run(child.get_text())
        elif element.name == 'ul':
            for li in element.find_all('li', recursive=False):
                document.add_paragraph(li.get_text(), style='List Bullet')
        elif element.name == 'ol':
            for li in element.find_all('li', recursive=False):
                document.add_paragraph(li.get_text(), style='List Number')
        elif element.name == 'blockquote':
            document.add_paragraph(element.get_text(), style='Quote')

    document.save(output_path)
    return output_path


def create_pdf_from_html(html_content: str, output_path: str):
    """
    Parses HTML content (generated from markdown) and creates a PDF file.
    """
    class PDF(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font("helvetica", "I", 8)
            self.cell(0, 10, f"Página {self.page_no()}", align="C")

    pdf = PDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    
    # fpdf2 write_html supports basic HTML tags correctly
    pdf.write_html(html_content)
    
    pdf.output(output_path)
    return output_path


def build_document(markdown_text: str, filename_base: str, output_dir: str):
    """
    Generates both .docx and .pdf files from a markdown string.
    Returns the file paths of the generated documents.
    """
    os.makedirs(output_dir, exist_ok=True)
    
    html_content = markdown.markdown(markdown_text)
    
    docx_path = os.path.join(output_dir, f"{filename_base}.docx")
    pdf_path = os.path.join(output_dir, f"{filename_base}.pdf")
    
    create_docx_from_html(html_content, docx_path)
    create_pdf_from_html(html_content, pdf_path)
    
    return docx_path, pdf_path
